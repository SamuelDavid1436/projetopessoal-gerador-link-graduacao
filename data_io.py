# -*- coding: utf-8 -*-
"""
data_io.py
==========
Funções para ler a lista de RAs (CSV ou Excel, apenas uma coluna) e para
gravar o resultado final, também em CSV e Excel.
"""
import csv as csv_modulo
import os
import re
from datetime import datetime

import pandas as pd

import config


def _detectar_separador_csv(caminho_arquivo: str):
    """
    Detecta o separador de um .csv de forma confiável. NÃO usa o
    sep=None/engine='python' do pandas pra isso — foi descoberto que esse
    modo erra sozinho (trunca valores!) em arquivos com uma única coluna
    de números e sem cabeçalho, que é exatamente o formato mais comum de
    planilha de RAs. csv.Sniffer() é bem mais confiável pra esse caso.

    Devolve o separador detectado, ou None se não achar nenhum de forma
    confiável — nesse caso, quem chama deve tratar a linha inteira como
    um campo só (arquivo de uma coluna sem delimitador nenhum).
    """
    with open(caminho_arquivo, "r", encoding="utf-8-sig", errors="ignore") as f:
        amostra = f.read(4096)
    if not amostra.strip():
        return None
    try:
        dialeto = csv_modulo.Sniffer().sniff(amostra, delimiters=",;\t")
        return dialeto.delimiter
    except csv_modulo.Error:
        return None  # sem delimitador confiável -- provavelmente uma coluna só


def extrair_ras_com_erro(caminho_resultado_csv: str) -> list:
    """
    Lê um arquivo resultado.csv de uma execução já concluída e devolve
    só os RAs das linhas cujo "Status da Consulta" indica erro (começa
    com "Erro") — usado pelo botão "Reprocessar erros", pra montar uma
    nova base só com quem falhou, sem precisar filtrar a planilha na mão.
    """
    if not os.path.isfile(caminho_resultado_csv):
        return []
    try:
        df = pd.read_csv(caminho_resultado_csv, sep=";", dtype=str, encoding="utf-8-sig")
    except (OSError, pd.errors.ParserError):
        return []
    if "RA" not in df.columns or "Status da Consulta" not in df.columns:
        return []

    status = df["Status da Consulta"].fillna("").astype(str).str.strip().str.lower()
    mascara_erro = status.str.startswith("erro")
    ras_com_erro = df.loc[mascara_erro, "RA"].dropna().astype(str).tolist()
    return ras_com_erro


def salvar_lista_ras(ras: list, caminho_destino: str):
    """Salva uma lista simples de RAs num .csv (uma coluna, com cabeçalho
    'RA') — usado pra gerar a base de reprocessamento de erros. Sem BOM
    (utf-8 puro): com BOM, o cabeçalho "RA" vira "\\ufeffRA" e o filtro de
    cabeçalho de ler_lista_ras não reconhece mais como cabeçalho."""
    os.makedirs(os.path.dirname(caminho_destino), exist_ok=True)
    df = pd.DataFrame({"RA": ras})
    df.to_csv(caminho_destino, index=False, encoding="utf-8")


def ler_lista_ras(caminho_arquivo: str) -> list:
    """
    Lê um arquivo CSV ou Excel contendo uma única coluna de RAs (com ou sem
    cabeçalho) e devolve uma lista de strings (RAs), sem valores vazios.

    IMPORTANTE: não remove RAs repetidos — se o mesmo RA aparecer mais de
    uma vez no arquivo importado, ele é processado todas as vezes que
    aparecer. Isso é proposital: a mesma base pode ser rodada várias vezes
    por semana (gerando link atualizado pra quem ainda não pagou), e nada
    aqui trava ou pula RA repetido, nem entre execuções nem dentro do
    mesmo arquivo.
    """
    extensao = os.path.splitext(caminho_arquivo)[1].lower()

    if extensao == ".csv":
        separador = _detectar_separador_csv(caminho_arquivo)
        if separador is None:
            # sem delimitador confiável -- trata a linha inteira como um
            # campo só (caso mais comum: uma coluna de RAs, sem cabeçalho)
            df = pd.read_csv(caminho_arquivo, sep="\x00", engine="python", header=None, dtype=str)
        else:
            df = pd.read_csv(caminho_arquivo, sep=separador, engine="python", header=None, dtype=str)
    elif extensao in (".xlsx", ".xls"):
        df = pd.read_excel(caminho_arquivo, header=None, dtype=str)
    else:
        raise ValueError(
            f"Formato de arquivo não suportado: {extensao}. Use .csv, .xlsx ou .xls."
        )

    # pega a primeira coluna, remove vazios e possível cabeçalho textual
    primeira_coluna = df.iloc[:, 0].dropna().astype(str).str.strip()
    ras = [ra for ra in primeira_coluna.tolist() if ra and ra.upper() not in ("RA", "MATRICULA", "MATRÍCULA")]

    return ras


def salvar_resultado(registros: list, pasta_saida: str = None, momento: datetime = None) -> tuple:
    """
    Salva a lista de registros (lista de dicionários, chaves = config.COLUNAS_SAIDA)
    numa SUBPASTA própria dentro da pasta de saída, nomeada com a data/hora
    da importação (`momento` — normalmente o início da execução). Assim,
    cada vez que uma base é importada e processada, os arquivos ficam
    isolados na pasta daquela execução específica, sem se misturar com
    execuções anteriores ou seguintes.

    Exemplo: duas importações no mesmo dia geram duas pastas diferentes:
        saida/2026-07-20_14-42-05/resultado.csv
        saida/2026-07-20_14-42-05/resultado.xlsx
        saida/2026-07-20_15-00-12/resultado.csv
        saida/2026-07-20_15-00-12/resultado.xlsx

    Retorna (caminho_csv, caminho_xlsx).
    """
    pasta_saida = pasta_saida or config.PASTA_SAIDA
    momento = momento or datetime.now()

    nome_pasta_execucao = momento.strftime("%Y-%m-%d_%H-%M-%S")
    pasta_execucao = os.path.join(pasta_saida, nome_pasta_execucao)
    os.makedirs(pasta_execucao, exist_ok=True)

    df = pd.DataFrame(registros, columns=config.COLUNAS_SAIDA)

    caminho_csv = os.path.join(pasta_execucao, "resultado.csv")
    caminho_xlsx = os.path.join(pasta_execucao, "resultado.xlsx")

    df.to_csv(caminho_csv, index=False, encoding="utf-8-sig", sep=";")
    df.to_excel(caminho_xlsx, index=False)

    return caminho_csv, caminho_xlsx


# Sub-colunas de cada bloco de mês no relatório largo. Cada bloco já tem o
# nome do mês no cabeçalho (ex: "Junho - Vencimento"); "Situacao Mensalidade"
# indica se aquele mês existe pra esse aluno (e o status da fatura).
_SUBCOLUNAS_POR_MES = ["Situacao Mensalidade", "Valor Pago", "Vencimento", "Link Pagamento"]

# Ordem final do relatório: identificação do aluno primeiro, depois a
# situação financeira, e só então um bloco fixo pra cada mês de
# config.MES_MINIMO_RELATORIO até config.MES_MAXIMO_RELATORIO.
_COLUNAS_IDENTIFICACAO = ["RA", "Nome", "CPF", "Telefone"]
_COLUNA_SITUACAO = "Situação"

# Colunas da base de disparo (nesta ordem). A última coluna ganha o nome do
# mês quando todos os alunos da base são do mesmo mês (ex: "Outubro - Link
# Pagamento"); com meses misturados, fica só "Link Pagamento".
_COLUNAS_BASE_DISPARO = ["RA", "CPF", "Nome", "Telefone", "MÊS", "Vencimento"]
_LARGURAS_BASE_DISPARO = [13, 14, 46, 16, 11, 12, 52]


def _meses_fixos_relatorio() -> list:
    """Lista travada de (mês, ano) — de config.MES_MINIMO_RELATORIO/
    ANO_MINIMO_RELATORIO até config.MES_MAXIMO_RELATORIO/ANO_MAXIMO_RELATORIO,
    inclusive nas duas pontas. Suporta virar o ano no meio (ex: Novembro/2026
    até Fevereiro/2027), caso um dia precise."""
    idx_ini = config.MESES.index(config.MES_MINIMO_RELATORIO)
    idx_fim = config.MESES.index(config.MES_MAXIMO_RELATORIO)
    ano_ini = int(config.ANO_MINIMO_RELATORIO)
    ano_fim = int(config.ANO_MAXIMO_RELATORIO)

    meses = []
    ano_atual, idx_atual = ano_ini, idx_ini
    while (ano_atual, idx_atual) <= (ano_fim, idx_fim):
        meses.append((config.MESES[idx_atual], str(ano_atual)))
        idx_atual += 1
        if idx_atual > 11:
            idx_atual = 0
            ano_atual += 1
    return meses


def _digitos_telefone(celular) -> str:
    """Devolve só os dígitos do celular, SEM o DDI 55 (DDD + número)."""
    digitos = re.sub(r"\D", "", str(celular or ""))
    while digitos.startswith("55") and len(digitos) > 11:
        digitos = digitos[2:]
    return digitos


def _telefone_formatado(celular) -> str:
    """'5511983224029' -> '(11)983224029' (formato do relatorio_meses)."""
    digitos = _digitos_telefone(celular)
    if len(digitos) < 10:
        return digitos
    return f"({digitos[:2]}){digitos[2:]}"


def _telefone_disparo(celular):
    """'(11)983224029' / '5511983224029' -> 5511983224029 (55 + DDD + número,
    como número inteiro). Devolve None se não houver telefone válido."""
    digitos = _digitos_telefone(celular)
    if len(digitos) not in (10, 11):
        return None
    return int("55" + digitos)


def _identificacao(registro: dict) -> dict:
    """Campos de identificação do aluno, já com os nomes de coluna do relatório."""
    return {
        "RA": registro.get("RA", ""),
        "Nome": str(registro.get("Nome", "") or ""),
        "CPF": str(registro.get("CPF", "") or ""),
        "Telefone": _telefone_formatado(registro.get("Celular", "")),
    }


def _pasta_da_execucao(pasta_saida, momento) -> str:
    pasta_saida = pasta_saida or config.PASTA_SAIDA
    momento = momento or datetime.now()
    pasta_execucao = os.path.join(pasta_saida, momento.strftime("%Y-%m-%d_%H-%M-%S"))
    os.makedirs(pasta_execucao, exist_ok=True)
    return pasta_execucao


def _gravar_csv_e_xlsx(df: pd.DataFrame, pasta_execucao: str, nome: str, larguras=None, colunas_numericas=()):
    """Grava df como <nome>.csv (UTF-8 com BOM, separado por ';') e
    <nome>.xlsx (cabeçalho em negrito). Tudo é texto no Excel, exceto as
    `colunas_numericas`."""
    caminho_csv = os.path.join(pasta_execucao, f"{nome}.csv")
    caminho_xlsx = os.path.join(pasta_execucao, f"{nome}.xlsx")
    df.to_csv(caminho_csv, index=False, encoding="utf-8-sig", sep=";")
    df.to_excel(caminho_xlsx, index=False)

    from openpyxl import load_workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = load_workbook(caminho_xlsx)
    ws = wb.active
    for celula in ws[1]:
        celula.font = Font(bold=True)
    for idx_col, coluna in enumerate(df.columns, start=1):
        letra = get_column_letter(idx_col)
        if larguras and idx_col <= len(larguras):
            ws.column_dimensions[letra].width = larguras[idx_col - 1]
        for linha in range(2, ws.max_row + 1):
            celula = ws.cell(row=linha, column=idx_col)
            if coluna in colunas_numericas:
                celula.number_format = "0"
            elif celula.value is not None:
                celula.number_format = "@"
    wb.save(caminho_xlsx)
    return caminho_csv, caminho_xlsx


def salvar_relatorio_meses(resultados: list, pasta_saida: str = None, momento: datetime = None) -> tuple:
    """
    Gera um arquivo SEPARADO (relatorio_meses.csv/.xlsx) — não altera o
    resultado.csv/.xlsx normal — com as parcelas de cada aluno, em formato
    LARGO: uma linha por aluno e um bloco de colunas TRAVADO pra cada mês
    entre config.MES_MINIMO_RELATORIO/ANO_MINIMO_RELATORIO e
    config.MES_MAXIMO_RELATORIO/ANO_MAXIMO_RELATORIO.

    Ordem das colunas: RA, Nome, CPF, Telefone, Situação e, pra cada mês, "<Mês> - Situacao Mensalidade",
    "<Mês> - Valor Pago", "<Mês> - Vencimento", "<Mês> - Link Pagamento".

    Cada mês em aberto (não Pago/Negociado) tem o link de pagamento gerado
    de verdade (ver crm_client.capturar_parcelas_a_partir_de); a data de
    vencimento vem da grade pra todo mês que tenha parcela.

    `resultados` é a mesma lista de registros da execução — cada um deve
    ter a chave "_parcelas_relatorio". Salva na MESMA pasta da execução
    (mesmo `momento`) que o resultado.csv/.xlsx normal.

    Retorna (caminho_csv, caminho_xlsx).
    """
    pasta_execucao = _pasta_da_execucao(pasta_saida, momento)
    meses_fixos = _meses_fixos_relatorio()  # [(mes, ano), ...] travado

    colunas_finais = list(_COLUNAS_IDENTIFICACAO) + [_COLUNA_SITUACAO]
    for mes, _ano in meses_fixos:
        for subcoluna in _SUBCOLUNAS_POR_MES:
            colunas_finais.append(f"{mes} - {subcoluna}")

    linhas = []
    for registro in resultados:
        linha = _identificacao(registro)
        linha[_COLUNA_SITUACAO] = registro.get("Situacao", "")
        parcelas = registro.get("_parcelas_relatorio", []) or []
        # indexa as parcelas desse aluno por (mês, ano) pra achar rápido
        parcelas_por_mes = {(p.get("Competencia"), str(p.get("Ano"))): p for p in parcelas}

        # Se a grade não foi consultada de verdade (erro/instabilidade), os
        # meses sem parcela são "Não consultado" -- nunca "Sem mensalidade".
        consultado = registro.get("_relatorio_consultado", True)
        sem_extratos = registro.get("Resultado da Consulta") == "SEM_EXTRATOS"

        algum_mes_preenchido = False
        for mes, ano in meses_fixos:
            dados_mes = parcelas_por_mes.get((mes, ano))
            vazio = {"Valor Pago": "", "Vencimento": "", "Link Pagamento": ""}
            if dados_mes is None and sem_extratos:
                situacao_mes, valores = "Sem extratos no CRM (conferir)", vazio
            elif dados_mes is None and not consultado:
                situacao_mes, valores = "Não consultado", vazio
            elif dados_mes is not None:
                # mostra o Status da Fatura junto, pra ficar claro POR QUE não
                # tem link (ex: "Pago"/"Negociado" não geram link, é esperado)
                status_fatura = str(dados_mes.get("Status da Fatura", "") or "").strip()
                situacao_mes = f"Com mensalidade ({status_fatura})" if status_fatura else "Com mensalidade"
                valores = {
                    "Valor Pago": dados_mes.get("Valor Pago", ""),
                    "Vencimento": dados_mes.get("Data Vencimento", ""),
                    "Link Pagamento": dados_mes.get("Link Pagamento", ""),
                }
                algum_mes_preenchido = True
            else:
                situacao_mes, valores = "Sem mensalidade", vazio
            linha[f"{mes} - Situacao Mensalidade"] = situacao_mes
            for subcoluna, valor in valores.items():
                linha[f"{mes} - {subcoluna}"] = valor

        # Linha toda vazia (sem Situação nem nenhum mês encontrado) --
        # sinaliza o motivo em vez de deixar tudo em branco.
        sem_situacao = not str(linha.get(_COLUNA_SITUACAO, "")).strip()
        if sem_situacao and not algum_mes_preenchido:
            status_consulta = str(registro.get("Status da Consulta", "") or "")
            eh_outro_erro = (
                status_consulta.lower().startswith("erro")
                and "não encontrado" not in status_consulta.lower()
            )
            if eh_outro_erro:
                linha[_COLUNA_SITUACAO] = f"Erro na consulta: {status_consulta[len('Erro:'):].strip()}"
            else:
                linha[_COLUNA_SITUACAO] = "RA não encontrado na base"

        linhas.append(linha)

    df = pd.DataFrame(linhas, columns=colunas_finais)
    return _gravar_csv_e_xlsx(df, pasta_execucao, "relatorio_meses")


def _data_para_ordenar(texto) -> tuple:
    """'07/10/2026' -> (2026, 10, 7). Data ilegível vira (0, 0, 0)."""
    m = re.match(r"\s*(\d{1,2})/(\d{1,2})/(\d{4})", str(texto or ""))
    return (int(m.group(3)), int(m.group(2)), int(m.group(1))) if m else (0, 0, 0)


def _parcelas_com_link(registro: dict) -> list:
    """Parcelas do aluno que TÊM link de pagamento gerado. Se o relatório de
    meses falhou pra esse RA, usa a parcela do fluxo principal (se tiver link)."""
    candidatas = []
    for p in registro.get("_parcelas_relatorio", []) or []:
        link = str(p.get("Link Pagamento", "") or "").strip()
        if link.lower().startswith("http"):
            candidatas.append({"mes": p.get("Competencia", ""), "ano": str(p.get("Ano", "")),
                               "vencimento": p.get("Data Vencimento", ""), "link": link})
    if not candidatas:
        link = str(registro.get("Link de Pagamento", "") or "").strip()
        if link.lower().startswith("http"):
            candidatas.append({"mes": registro.get("Competencia", ""), "ano": str(registro.get("Ano", "")),
                               "vencimento": registro.get("Vencimento", ""), "link": link})
    return candidatas


def salvar_base_disparo(resultados: list, pasta_saida: str = None, momento: datetime = None) -> tuple:
    """
    Gera a base de disparo (base_disparo.csv/.xlsx): UMA linha por aluno,
    com o link da parcela de vencimento mais recente entre as que tiveram
    link gerado. Alunos sem nenhum link ficam de fora.

    Colunas: RA, CPF, Nome, Telefone (55 + DDD + número), MÊS, Vencimento e o link ("<Mês> - Link Pagamento",
    ou só "Link Pagamento" se a base tiver meses misturados).

    Retorna (caminho_csv, caminho_xlsx).
    """
    pasta_execucao = _pasta_da_execucao(pasta_saida, momento)

    linhas, meses_usados = [], []
    for registro in resultados:
        candidatas = _parcelas_com_link(registro)
        if not candidatas:
            continue
        escolhida = max(
            candidatas,
            key=lambda c: (_data_para_ordenar(c["vencimento"]), int(c["ano"]) if c["ano"].isdigit() else 0,
                           config.MESES.index(c["mes"]) if c["mes"] in config.MESES else -1),
        )
        ident = _identificacao(registro)
        meses_usados.append(escolhida["mes"])
        linhas.append({
            "RA": ident["RA"], "CPF": ident["CPF"], "Nome": ident["Nome"],
            "Telefone": _telefone_disparo(registro.get("Celular", "")),
            "MÊS": escolhida["mes"], "Vencimento": escolhida["vencimento"], "_link": escolhida["link"],
        })

    unicos = {m for m in meses_usados if m}
    nome_coluna_link = f"{next(iter(unicos))} - Link Pagamento" if len(unicos) == 1 else "Link Pagamento"
    df = pd.DataFrame(linhas, columns=_COLUNAS_BASE_DISPARO + ["_link"]).rename(columns={"_link": nome_coluna_link})
    df["Telefone"] = df["Telefone"].astype("Int64")
    return _gravar_csv_e_xlsx(df, pasta_execucao, "base_disparo",
                              larguras=_LARGURAS_BASE_DISPARO, colunas_numericas=("Telefone",))
